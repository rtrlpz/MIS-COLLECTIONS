from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
import pandas as pd

DATA = Path(__file__).resolve().parents[1] / "data"

OUT = Path(__file__).resolve().parents[1] / "daily_mis_auto.xlsx"
OUT.parent.mkdir(parents=True, exist_ok=True)

wb = Workbook()
ws = wb.active
ws.title = "Daily MIS Auto"

ws["A1"] = "Collections Daily MIS"
ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="262A76")

headers = ["Date", "Contacts", "Connects", "RPCs", "Promises", "Payments"]
for c, h in enumerate(headers, start=2):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = Font(bold=True, color="262A76", name="Calibri")
    cell.alignment = Alignment(horizontal="center")

ws["B2"] = "2025-06-30"
ws["A2"] = '="Report for "&TEXT(B2, "yyyy-mm-dd")'
ws["A2"].font = Font(italic=True)

ws.freeze_panes = "A5"
for c in range(2, 8):
    ws.column_dimensions[get_column_letter(c)].width = 14

wb.save(OUT)
print(f"Saved to {OUT}")

### LOAD DATA ###
interactions = pd.read_csv(DATA / "raw" / "interactions_all.csv",
                           parse_dates=["interaction_date"])
ptp = pd.read_csv(DATA / "raw"/ "promises_data.csv",
                  parse_dates=["ptp_date"])
payments = pd.read_csv(DATA / "raw" / "payments_data.csv",
                       parse_dates=["payment_date"])

it_day = interactions["interaction_date"].dt.normalize()
ptp_day = ptp["ptp_date"].dt.normalize()
pay_day = payments["payment_date"].dt.normalize()

dates = pd.DatetimeIndex(sorted(set(it_day) | set(ptp_day) | set(pay_day)))
daily = pd.DataFrame(index=dates)
daily.index.name = "Date"

daily["Contacts"] = it_day.value_counts().reindex(dates, fill_value=0)
daily["Connects"] = interactions.groupby(it_day)["calls_connected"].sum().reindex(dates, fill_value=0)
daily["RPCs"] = interactions.groupby(it_day)["rpc_flag"].sum().reindex(dates, fill_value=0)
daily["Promises"] = ptp_day.value_counts().reindex(dates, fill_value=0)
daily["Payments"] = pay_day.value_counts().reindex(dates, fill_value=0)

daily = daily.astype(int).reset_index()

wb = load_workbook(OUT)
ws = wb["Daily MIS Auto"]
for r, row in daily.iterrows():
    c = ws.cell(row=5 + r, column=2, value=row["Date"].to_pydatetime())
    c.number_format = "yyyy-mm-dd"
    for col_idx, col in enumerate(["Contacts", "Connects", "RPCs", "Promises", "Payments"], start=3):
        cell = ws.cell(row=5 + r, column=col_idx, value=int(row[col]))
        cell.number_format = "#,##0"
wb.save(OUT)
print(f"Daily data loaded into: {OUT}")

### XLOOKUP AND VLOOKUP ###
wb = load_workbook(OUT)

emp = pd.read_csv(DATA / "raw" / "employees_data.csv", dtype="string")[["agent_id", "agent_name"]]
if "AgentLookup" in wb.sheetnames:
    del wb["AgentLookup"]
ws_lu = wb.create_sheet("AgentLookup")
ws_lu.append(["AgentID", "Name"])
for agent_id, agent_name in emp.itertuples(index=False, name=None):
    ws_lu.append([agent_id, agent_name])
n_last = len(emp) + 1                      # 89 -> data rows 2..89

hdr_font = Font(bold=True, color="262A76", name="Calibri")
ctr = Alignment(horizontal="center")

for col_idx, title in ((9, "AgentID"), (10, "AgentName"), (11, "LookupType")):
    c = ws.cell(row=4, column=col_idx, value=title)
    c.font = hdr_font
    c.alignment = ctr

ws.cell(row=5, column=9, value="EID-001").alignment = ctr
ws.cell(row=6, column=9, value="EID-042").alignment = ctr

ws.cell(row=5, column=10, value=(
    f'=IFNA(_xlfn.XLOOKUP(I5,AgentLookup!$A$2:$A${n_last},AgentLookup!$B$2:$B${n_last}),"—")'))
ws.cell(row=6, column=10, value=(
    f'=IFNA(VLOOKUP(I6,AgentLookup!$A$2:$B${n_last},2,FALSE),"—")'))

ws.cell(row=5, column=11, value="XLOOKUP").alignment = ctr
ws.cell(row=6, column=11, value="VLOOKUP").alignment = ctr

for col in range(9, 12):
    ws.column_dimensions[get_column_letter(col)].width = 16

wb.save(OUT)
print(f"Lookup demo written to: {OUT}")